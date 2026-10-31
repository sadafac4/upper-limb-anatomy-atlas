"""Build a .pptx deck from a YAML slide spec, then render PDF + PNG previews.

Usage:
  python pipeline/build_deck.py topografik/slides/pilot.yaml topografik/out/pilot.pptx

Spec format (YAML list):
  - title: "Plexus brachialis — fasciculi"
    image: topografik/images/<file>          # must have a row in images/PROVENANCE.csv
    caption: "Kaynak: ..., Lisans: CC BY 4.0"
    bullets: ["...Türkçe metin, Latince terimler..."]
    labels:                                   # annotations; x,y are 0-1 relative to image
      - {text: "Fasciculus lateralis", x: 0.42, y: 0.31, tx: 0.85, ty: 0.20}
    citation: "Yazar A, et al. Dergi. 2020;12:34-40. doi:..."

Rules enforced: every image must exist and be listed in PROVENANCE.csv with a license.
"""
import csv, subprocess, sys
from pathlib import Path
import yaml
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor

W, H = Emu(12192000), Emu(6858000)  # 16:9
ROOT = Path(__file__).resolve().parents[1]


def load_provenance(img: Path):
    prov = img.parent / "PROVENANCE.csv"
    if not prov.exists():
        sys.exit(f"Missing {prov}")
    with prov.open(newline="", encoding="utf-8") as f:
        rows = {r["file"]: r for r in csv.DictReader(f)}
    row = rows.get(img.name)
    if not row or not row.get("license"):
        sys.exit(f"No provenance/license row for {img.name} in {prov}")
    return row


def add_slide(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(400000), Emu(200000), W - Emu(800000), Emu(700000))
    tb.text_frame.text = s["title"]
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True

    box_l, box_t, box_w, box_h = Emu(400000), Emu(1000000), Emu(7000000), Emu(5200000)
    if s.get("image"):
        img = ROOT / s["image"]
        load_provenance(img)
        iw, ih = Image.open(img).size
        scale = min(box_w / iw, box_h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        pic = slide.shapes.add_picture(str(img), box_l, box_t, pw, ph)
        for lab in s.get("labels", []):
            px, py = pic.left + int(lab["x"] * pw), pic.top + int(lab["y"] * ph)
            tx, ty = pic.left + int(lab["tx"] * pw), pic.top + int(lab["ty"] * ph)
            ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, tx, ty, px, py)
            ln.line.color.rgb = RGBColor(0xFF, 0xD0, 0x00)
            ln.line.width = Pt(1.5)
            t = slide.shapes.add_textbox(tx, ty - Emu(150000), Emu(2200000), Emu(300000))
            t.text_frame.text = lab["text"]
            r = t.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.italic = Pt(12), True
            r.font.color.rgb = RGBColor(0xFF, 0xD0, 0x00)

    body = slide.shapes.add_textbox(Emu(7600000), Emu(1000000), Emu(4200000), Emu(5000000)).text_frame
    body.word_wrap = True
    for i, b in enumerate(s.get("bullets", [])):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = "• " + b
        p.runs[0].font.size = Pt(14)

    foot = " | ".join(x for x in (s.get("caption"), s.get("citation")) if x)
    if foot:
        f = slide.shapes.add_textbox(Emu(400000), H - Emu(550000), W - Emu(800000), Emu(400000)).text_frame
        f.word_wrap = True
        f.text = foot
        f.paragraphs[0].runs[0].font.size = Pt(9)


def render(pptx_path: Path):
    out = pptx_path.parent
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx_path)], check=True)
    pdf = pptx_path.with_suffix(".pdf")
    prev = out / "preview"
    prev.mkdir(exist_ok=True)
    subprocess.run(["pdftoppm", "-png", "-r", "60", str(pdf), str(prev / pptx_path.stem)], check=True)
    return pdf


def main(spec, out):
    slides = yaml.safe_load(Path(spec).read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for s in slides:
        add_slide(prs, s)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print("saved", out, "->", render(out))


if __name__ == "__main__":
    main(*sys.argv[1:3])
