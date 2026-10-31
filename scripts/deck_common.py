#!/usr/bin/env python3
"""Shared python-pptx helpers for the Turkish upper-limb anatomy decks.

Annotations are native PowerPoint shapes (editable): a small marker on the structure, a leader
line and a text box placed in a label column outside the image. Point coordinates are given in
SOURCE-IMAGE PIXELS and converted to slide EMU from the actual placed picture geometry, so labels
stay correct whatever the scaling.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
YELLOW = RGBColor(0xFF, 0xD2, 0x00)
FONT = "Calibri"

def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs

def _tf_style(tf, size, color=NAVY, bold=False, align=PP_ALIGN.LEFT, italic=False):
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
            r.font.color.rgb, r.font.name = color, FONT

def text(slide, x, y, w, h, s, size=14, color=NAVY, bold=False, align=PP_ALIGN.LEFT, italic=False,
         anchor=MSO_ANCHOR.TOP, fill=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _add_runs(p, line)
    _tf_style(tf, size, color, bold, align, italic)
    if fill is not None:
        tb.fill.solid(); tb.fill.fore_color.rgb = fill
    return tb

def _add_runs(p, line):
    """Support **bold** and ~italic~ fragments inside a line."""
    import re
    parts = re.split(r"(\*\*[^*]+\*\*|~[^~]+~)", line)
    for part in parts:
        if not part: continue
        r = p.add_run()
        if part.startswith("**"): r.text = part[2:-2]; r._bold = True
        elif part.startswith("~") and part.endswith("~") and len(part) > 2: r.text = part[1:-1]; r._ital = True
        else: r.text = part

def fix_runs(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if getattr(r, "_bold", False): r.font.bold = True
                if getattr(r, "_ital", False): r.font.italic = True

def bullets(slide, x, y, w, h, items, size=13, color=NAVY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _add_runs(p, "• " + it)
        p.space_after = Pt(3)
    _tf_style(tf, size, color)
    return tb

def title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.85))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text(slide, 0.3, 0.08, 12.7, 0.5, title, size=24, color=WHITE, bold=True)
    if subtitle:
        text(slide, 0.3, 0.5, 12.7, 0.32, subtitle, size=13, color=RGBColor(0xC8, 0xD1, 0xE6), italic=True)

def footer(slide, s, page=None):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.3), Inches(7.0), Inches(13.03), Inches(7.0))
    line.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB); line.line.width = Pt(0.75)
    text(slide, 0.3, 7.02, 12.0, 0.45, s, size=8.5, color=GREY)
    if page is not None:
        text(slide, 12.4, 7.05, 0.65, 0.3, str(page), size=10, color=GREY, align=PP_ALIGN.RIGHT)

def picture_fit(slide, path, x, y, w, h, align="center"):
    """Place image inside the box keeping aspect; returns geometry dict for coordinate mapping."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = x + (w - pw) / 2 if align == "center" else (x if align == "left" else x + w - pw)
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return dict(x=px, y=py, w=pw, h=ph, iw=iw, ih=ih, pic=pic)

def img2slide(geo, pt):
    return geo["x"] + pt[0] / geo["iw"] * geo["w"], geo["y"] + pt[1] / geo["ih"] * geo["h"]

def labels(slide, geo, items, col_x, col_w, side, y_min, y_max, size=11.5, color=NAVY,
           line_color=YELLOW, box_h=0.36):
    """items: list of (label_text, (px, py)) in image pixels. Labels are stacked in a column at
    col_x (inches) on the given side ('L' → text right-aligned, line leaves from its right edge)."""
    pts = [(t, img2slide(geo, p)) for t, p in items]
    pts.sort(key=lambda a: a[1][1])
    n = len(pts)
    # initial y = point y, then resolve overlaps top-down / bottom-up
    ys = [min(max(p[1][1] - box_h / 2, y_min), y_max - box_h) for p in pts]
    gap = box_h + 0.06
    for i in range(1, n):
        ys[i] = max(ys[i], ys[i - 1] + gap)
    if n and ys[-1] > y_max - box_h:
        ys[-1] = y_max - box_h
        for i in range(n - 2, -1, -1):
            ys[i] = min(ys[i], ys[i + 1] - gap)
    shapes = []
    for (t, (sx, sy)), ty in zip(pts, ys):
        tb = text(slide, col_x, ty, col_w, box_h, t, size=size, color=color, bold=True,
                  align=PP_ALIGN.RIGHT if side == "L" else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        lx = col_x + col_w if side == "L" else col_x
        ly = ty + box_h / 2
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(lx), Inches(ly), Inches(sx), Inches(sy))
        ln.line.color.rgb = line_color; ln.line.width = Pt(1.5)
        r = 0.045
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx - r), Inches(sy - r), Inches(2 * r), Inches(2 * r))
        dot.fill.solid(); dot.fill.fore_color.rgb = line_color; dot.line.color.rgb = NAVY; dot.line.width = Pt(0.5)
        shapes.append((tb, ln, dot))
    return shapes

def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s
