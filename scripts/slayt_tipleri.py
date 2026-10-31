# -*- coding: utf-8 -*-
"""Uzun (bölümlü) sunumlar için ortak slayt tipleri. deck_common üzerine kuruludur."""
import pathlib
from PIL import Image
from pptx.util import Inches, Pt
from deck_common import *
import kaynak_kayit as K

ROOT = pathlib.Path(__file__).resolve().parent.parent
VHP = ROOT / "kesitsel" / "images" / "vhp"
ETIKET_RENK = {"KONU ANLATIMI": RGBColor(0x1F, 0x5F, 0x9E), "KLİNİK BAĞLANTI": RGBColor(0xC2, 0x6A, 0x00),
               "DİSEKSİYON": RGBColor(0x2E, 0x7D, 0x32), "GÖRÜNTÜLEME": RGBColor(0x5E, 0x35, 0xB1),
               "VARYASYON": ACCENT, "LİTERATÜR TARAMASI": ACCENT, "KESİT": RGBColor(0x00, 0x69, 0x7C),
               "ŞEMA": RGBColor(0x45, 0x5A, 0x64), "TABLO": RGBColor(0x45, 0x5A, 0x64)}

def yeni_slayt(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def etiket(s, ad):
    text(s, 10.75, 0.9, 2.35, 0.28, ad, size=9.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         fill=ETIKET_RENK.get(ad, NAVY), anchor=MSO_ANCHOR.MIDDLE)

def bolum_seridi(s, bolum):
    if bolum:
        text(s, 0.3, 0.9, 4.0, 0.28, bolum, size=9.5, color=GREY, bold=True)

# ---------- görsel referans çözümleme ----------
def cozumle(ref):
    """ref → (yol, kısa altyazı, tam atıf veya None)"""
    tip = ref[0]
    if tip == "pmc":
        _, pmc, key = ref[:3]
        f = K.sekil(pmc, key)
        alt = ref[3] if len(ref) > 3 else f"{f['label']} – {K.atif(pmc, kisa=True)}"
        return ROOT / f["local_path"], alt, K.atif(pmc)
    if tip == "vhp":
        return VHP / ref[1], ref[2], None
    if tip == "vhp_crop":
        return VHP / f"cryo_{ref[1]}_R.jpg", ref[2] if len(ref) > 2 else f"VHP-M kriyokesit {ref[1]}", None
    raise ValueError(ref)

VHP_ATIF = ("Görüntü: NLM Visible Human Project – Visible Human Male (NCI Imaging Data Commons, doi:10.5281/zenodo.12690049; "
            "NLM Terms and Conditions 2019). İşleme/kırpma: bu proje.")

def gorselleri_yerlestir(s, refs, X, Y, W, H, yatay=None):
    """Görselleri kutuya en-boy oranına göre yan yana veya üst üste yerleştirir; atıf listesi döner."""
    items = [cozumle(r) for r in refs]
    sizes = [Image.open(p).size for p, _, _ in items]
    ar = [w / h for w, h in sizes]
    n = len(items)
    if n == 1:
        boxes = [(X, Y, W, H - 0.25)]
    else:
        if yatay is None:
            # hangi yerleşim daha büyük alan verir?
            hs = min(H - 0.3, (W - 0.15 * (n - 1)) / sum(ar))
            vs = min(W, (H - 0.3 * n) / sum(1 / a for a in ar))
            yatay = hs * sum(ar) * hs >= vs * vs * sum(1 / a for a in ar) * 0.9
        if yatay:
            h = min(H - 0.3, (W - 0.15 * (n - 1)) / sum(ar)); x = X
            boxes = []
            for a in ar:
                boxes.append((x, Y, a * h, h)); x += a * h + 0.15
        else:
            w = min(W, (H - 0.3 * n) / sum(1 / a for a in ar)); y = Y
            boxes = []
            for a in ar:
                boxes.append((X + (W - w) / 2, y, w, w / a)); y += w / a + 0.3
    atiflar = []
    for (p, alt, atf), b in zip(items, boxes):
        g = picture_fit(s, p, *b)
        text(s, g["x"] - 0.1, g["y"] + g["h"], g["w"] + 0.2, 0.25, alt, size=8, color=GREY, italic=True, align=PP_ALIGN.CENTER)
        atiflar.append(atf or VHP_ATIF)
    return list(dict.fromkeys(atiflar))

def dipnot(s, atiflar, sayfa):
    footer(s, "  •  ".join(atiflar) if atiflar else "", sayfa)

# ---------- slayt tipleri ----------
def kapak(prs, baslik, alt, satirlar, refs):
    s = yeni_slayt(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    text(s, 0.7, 1.1, 7.4, 1.5, baslik, size=36, color=WHITE, bold=True)
    text(s, 0.7, 2.75, 7.4, 1.0, alt, size=17, color=RGBColor(0xC8, 0xD1, 0xE6))
    text(s, 0.7, 4.1, 7.4, 2.4, satirlar, size=13.5, color=WHITE)
    x = 8.5
    for i, r in enumerate(refs[:4]):
        p, _, _ = cozumle(r)
        picture_fit(s, p, 8.4 + (i % 2) * 2.35, 0.45 + (i // 2) * 3.45, 2.25, 3.3)
    text(s, 0.7, 6.75, 7.6, 0.4, "Sunum dili Türkçe; anatomik terimler Terminologia Anatomica (Latince).", size=11, color=RGBColor(0xC8, 0xD1, 0xE6), italic=True)
    fix_runs(s)

def icindekiler(prs, bolumler, ek=None, sayfa=None):
    s = yeni_slayt(prs); title_bar(s, "İçindekiler", "Her bölüm: konu anlatımı → klinik bağlantılar → varyasyon literatür taraması → seçilmiş çalışmalar")
    y = 1.15
    for i, b in enumerate(bolumler):
        kutu = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5 + i * 3.15), Inches(y), Inches(2.95), Inches(5.4))
        kutu.fill.solid(); kutu.fill.fore_color.rgb = LIGHT; kutu.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        text(s, 0.6 + i * 3.15, y + 0.1, 2.75, 0.5, f"{i + 1}. {b['ad']}", size=20, color=ACCENT, bold=True)
        text(s, 0.6 + i * 3.15, y + 0.65, 2.75, 0.45, b["alt"], size=10.5, color=GREY, italic=True)
        bullets(s, 0.6 + i * 3.15, y + 1.2, 2.75, 4.1, b["konular"], size=11)
    if ek:
        text(s, 0.5, 6.62, 12.3, 0.35, ek, size=11, color=GREY, italic=True)
    footer(s, "", sayfa); fix_runs(s)

def bolum_kapak(prs, no, b, refs, sayfa):
    s = yeni_slayt(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(7.2), prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    text(s, 0.6, 0.9, 6.3, 0.6, f"BÖLÜM {no}", size=18, color=RGBColor(0xC8, 0xD1, 0xE6), bold=True)
    text(s, 0.6, 1.45, 6.3, 1.0, b["ad"], size=48, color=WHITE, bold=True)
    text(s, 0.6, 2.6, 6.3, 0.9, [b["baslik"], b["alt"]], size=16, color=RGBColor(0xC8, 0xD1, 0xE6))
    bullets(s, 0.6, 3.75, 6.3, 3.2, b["konular"], size=13, color=WHITE)
    atif = gorselleri_yerlestir(s, refs, 7.45, 0.4, 5.6, 6.3)
    text(s, 7.45, 6.9, 5.6, 0.55, "  •  ".join(atif), size=7, color=GREY)
    text(s, 12.4, 7.05, 0.65, 0.3, str(sayfa), size=10, color=GREY, align=PP_ALIGN.RIGHT); fix_runs(s)

def anlatim(prs, d, sayfa, bolum=None):
    s = yeni_slayt(prs)
    title_bar(s, d["baslik"], d.get("alt")); bolum_seridi(s, bolum)
    etiket(s, "KLİNİK BAĞLANTI" if d["tip"] == "klinik" else d.get("etiket", "KONU ANLATIMI"))
    figs = d.get("figs") or []
    atif = []
    if figs:
        bullets(s, 0.3, 1.3, 7.2, 5.6, d["metin"], size=d.get("boyut", 13))
        atif = gorselleri_yerlestir(s, figs, 7.7, 1.3, 5.35, 5.55)
    else:
        bullets(s, 0.4, 1.35, 12.5, 5.5, d["metin"], size=d.get("boyut", 16.5))
    dipnot(s, atif, sayfa); fix_runs(s)
    if d.get("not"): notes(s, d["not"])
    return s

def tablo(prs, d, sayfa, bolum=None):
    s = yeni_slayt(prs)
    title_bar(s, d["baslik"], d.get("alt")); bolum_seridi(s, bolum); etiket(s, "TABLO")
    rows, cols = len(d["satirlar"]) + 1, len(d["sutunlar"])
    H = min(5.3, 0.45 + d.get("satir_h", 0.75) * (rows - 1))
    gt = s.shapes.add_table(rows, cols, Inches(0.3), Inches(1.3), Inches(12.73), Inches(H)).table
    gen = d.get("genislik") or [1] * cols
    tot = sum(gen)
    for j, g in enumerate(gen):
        gt.columns[j].width = Inches(12.73 * g / tot)
    for j, h in enumerate(d["sutunlar"]):
        c = gt.cell(0, j); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        for p in c.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    for i, row in enumerate(d["satirlar"], 1):
        for j, v in enumerate(row):
            c = gt.cell(i, j); c.text = v
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(d.get("boyut", 11)); r.font.color.rgb = NAVY; r.font.name = FONT
                    r.font.bold = (j == 0)
    if d.get("dip"):
        text(s, 0.3, 6.6, 12.7, 0.35, d["dip"], size=10, color=GREY, italic=True)
    footer(s, d.get("kaynak", "Kaynak: standart anatomi bilgisi (Terminologia Anatomica)."), sayfa); fix_runs(s)
    return s

def _kutu(s, x, y, w, h, t, renk, yazi=WHITE, boyut=11):
    k = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    k.fill.solid(); k.fill.fore_color.rgb = renk; k.line.color.rgb = WHITE
    tf = k.text_frame; tf.margin_left = tf.margin_right = Inches(0.03); tf.word_wrap = True
    tf.text = t
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs: r.font.size = Pt(boyut); r.font.bold = True; r.font.color.rgb = yazi; r.font.name = FONT
    return (x, y, w, h)

def _bag(s, a, b, renk=GREY, kalin=1.5):
    ax, ay, aw, ah = a; bx, by, bw, bh = b
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ax + aw), Inches(ay + ah / 2), Inches(bx), Inches(by + bh / 2))
    ln.line.color.rgb = renk; ln.line.width = Pt(kalin)

def sema_plexus(prs, d, sayfa, bolum=None):
    s = yeni_slayt(prs)
    title_bar(s, d["baslik"], d.get("alt")); bolum_seridi(s, bolum); etiket(s, "ŞEMA")
    SAR = RGBColor(0xE0, 0xA1, 0x00); YES = RGBColor(0x2E, 0x7D, 0x32); MAV = RGBColor(0x1F, 0x5F, 0x9E)
    xs = [0.4, 1.85, 3.35, 4.95, 6.6]; w = 1.25; h = 0.42
    for x, t in zip(xs, ["Radices", "Trunci", "Divisiones", "Fasciculi", "Rami terminales"]):
        text(s, x - 0.05, 1.2, w + 0.3, 0.3, t, size=11, color=GREY, bold=True, align=PP_ALIGN.CENTER)
    R = {k: _kutu(s, xs[0], 1.6 + i * 0.95, w, h, k, MAV) for i, k in enumerate(["C5", "C6", "C7", "C8", "T1"])}
    TR = {"sup": _kutu(s, xs[1], 2.05, w, h, "T. superior", MAV), "med": _kutu(s, xs[1], 3.5, w, h, "T. medius", MAV),
          "inf": _kutu(s, xs[1], 4.95, w, h, "T. inferior", MAV)}
    for r, t in [("C5", "sup"), ("C6", "sup"), ("C7", "med"), ("C8", "inf"), ("T1", "inf")]: _bag(s, R[r], TR[t])
    DV = {}
    for i, (t, y) in enumerate([("sup", 1.75), ("med", 3.2), ("inf", 4.65)]):
        DV[t + "A"] = _kutu(s, xs[2], y, w, 0.36, "ön", SAR, NAVY, 10); DV[t + "P"] = _kutu(s, xs[2], y + 0.5, w, 0.36, "arka", YES, WHITE, 10)
        _bag(s, TR[t], DV[t + "A"]); _bag(s, TR[t], DV[t + "P"])
    F = {"lat": _kutu(s, xs[3], 2.0, w, h, "F. lateralis", SAR, NAVY), "post": _kutu(s, xs[3], 3.55, w, h, "F. posterior", YES),
         "med": _kutu(s, xs[3], 5.1, w, h, "F. medialis", SAR, NAVY)}
    for a, f in [("supA", "lat"), ("medA", "lat"), ("supP", "post"), ("medP", "post"), ("infP", "post"), ("infA", "med")]: _bag(s, DV[a], F[f])
    TT = {"mc": _kutu(s, xs[4], 1.55, w + 0.15, h, "N. musculocutaneus", SAR, NAVY, 10), "ax": _kutu(s, xs[4], 2.85, w + 0.15, h, "N. axillaris", YES, WHITE, 10),
          "rad": _kutu(s, xs[4], 3.55, w + 0.15, h, "N. radialis", YES, WHITE, 10), "med": _kutu(s, xs[4], 4.45, w + 0.15, h, "N. medianus", SAR, NAVY, 10),
          "uln": _kutu(s, xs[4], 5.5, w + 0.15, h, "N. ulnaris", SAR, NAVY, 10)}
    for f, t in [("lat", "mc"), ("lat", "med"), ("post", "ax"), ("post", "rad"), ("med", "med"), ("med", "uln")]: _bag(s, F[f], TT[t])
    text(s, 0.4, 6.35, 7.6, 0.5, "Sarı: ön divisiolar → fleksör sinirler; yeşil: arka divisiolar → ekstensör sinirler. Şema basitleştirilmiştir (kollateral dallar sağdaki metinde).", size=9.5, color=GREY, italic=True)
    bullets(s, 8.25, 1.3, 4.8, 5.5, d["metin"], size=12)
    footer(s, "Şema: bu projede standart anatomi bilgisinden çizilmiştir (gerçek görüntü değildir).", sayfa); fix_runs(s)

def sema_arter(prs, d, sayfa, bolum=None):
    s = yeni_slayt(prs)
    title_bar(s, d["baslik"], d.get("alt")); bolum_seridi(s, bolum); etiket(s, "ŞEMA")
    KIR = RGBColor(0xB7, 0x1C, 0x1C)
    zinc = [("A. subclavia", "kosta I lateral kenarı"), ("A. axillaris", "m. teres major alt kenarı"), ("A. brachialis", "fossa cubiti (collum radii)")]
    y = 1.35; prev = None
    for ad, sinir in zinc:
        b = _kutu(s, 0.6, y, 2.4, 0.5, ad, KIR)
        if prev:
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.8), Inches(prev[1] + 0.5), Inches(1.8), Inches(y)); ln.line.color.rgb = KIR; ln.line.width = Pt(2.5)
        text(s, 3.1, y + 0.55, 3.9, 0.3, f"↓ sınır: {sinir}", size=10, color=GREY, italic=True)
        prev = b; y += 1.05
    dal = {"A. axillaris": "1: a. thoracica superior • 2: a. thoracoacromialis, a. thoracica lateralis • 3: a. subscapularis, aa. circumflexae humeri ant./post.",
           "A. brachialis": "a. profunda brachii • aa. collaterales ulnares sup./inf. • a. nutricia humeri"}
    text(s, 3.1, 2.4, 4.4, 0.55, dal["A. axillaris"], size=9.5, color=NAVY)
    text(s, 3.1, 3.45, 4.4, 0.5, dal["A. brachialis"], size=9.5, color=NAVY)
    r = _kutu(s, 0.2, y + 0.1, 1.7, 0.5, "A. radialis", KIR); u = _kutu(s, 2.1, y + 0.1, 1.7, 0.5, "A. ulnaris", KIR)
    for b in (r, u):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.8), Inches(prev[1] + 0.5), Inches(b[0] + b[2] / 2), Inches(b[1])); ln.line.color.rgb = KIR; ln.line.width = Pt(2.5)
    text(s, 3.9, y + 0.15, 3.6, 0.5, "a. ulnaris → a. interossea communis (ant./post.); rekürren dallar → dirsek anastomozu", size=9.5, color=NAVY)
    y2 = y + 1.15
    p = _kutu(s, 0.2, y2, 1.7, 0.6, "Arcus palmaris profundus", RGBColor(0x8E, 0x24, 0xAA), boyut=9.5)
    q = _kutu(s, 2.1, y2, 1.7, 0.6, "Arcus palmaris superficialis", RGBColor(0x8E, 0x24, 0xAA), boyut=9.5)
    for a, b in ((r, p), (u, q)):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0] + a[2] / 2), Inches(a[1] + 0.5), Inches(b[0] + b[2] / 2), Inches(b[1])); ln.line.color.rgb = KIR; ln.line.width = Pt(2.5)
    bullets(s, 7.7, 1.3, 5.35, 5.5, d["metin"], size=12)
    footer(s, "Şema: bu projede standart anatomi bilgisinden çizilmiştir (gerçek görüntü değildir).", sayfa); fix_runs(s)

def gorsel(prs, d, sayfa, bolum=None):
    """Görsel ağırlıklı slayt: solda büyük görsel(ler), sağda lejant + önemli noktalar."""
    s = yeni_slayt(prs)
    title_bar(s, d["baslik"], d.get("alt")); bolum_seridi(s, bolum); etiket(s, d.get("etiket", "DİSEKSİYON"))
    atif = gorselleri_yerlestir(s, d["figs"], 0.3, 1.3, 7.75, 5.55, d.get("yatay"))
    y = 1.3
    if d.get("lejant"):
        text(s, 8.3, y, 4.8, 0.3, "Şekildeki etiketler", size=12, color=ACCENT, bold=True)
        tb = bullets(s, 8.3, y + 0.3, 4.8, 2.6, d["lejant"], size=d.get("lboyut", 10.5)); y = 3.75 if d.get("noktalar") else y
    if d.get("noktalar"):
        if not d.get("lejant"): y = 1.3
        text(s, 8.3, y, 4.8, 0.3, d.get("nbaslik", "Önemli noktalar"), size=12, color=ACCENT, bold=True)
        bullets(s, 8.3, y + 0.3, 4.8, 6.85 - y - 0.3, d["noktalar"], size=d.get("nboyut", 10.5))
    if any(r[0] == "pmc" and K.nc_mi(r[1]) for r in d["figs"]):
        text(s, 8.3, 6.62, 4.8, 0.3, "Lisans: CC BY-NC – yalnızca ticari olmayan eğitim amaçlı kullanım.", size=8.5, color=ACCENT, italic=True)
    dipnot(s, atif, sayfa); fix_runs(s)
    notes(s, "\n".join(K.atif(r[1]) + "\nAltyazı: " + K.sekil(r[1], r[2])["caption"] for r in d["figs"] if r[0] == "pmc"))
    return s

def lit_tablo(prs, baslik, alt, satirlar, sayfa, bolum=None):
    """Varyasyon literatür taraması tablosu. satır: [varyasyon, çalışma, örneklem/yöntem, sıklık/bulgu, klinik önem]"""
    d = dict(baslik=baslik, alt=alt, sutunlar=["Varyasyon", "Çalışma", "Örneklem / yöntem", "Sıklık / ana bulgu", "Klinik önem"],
             satirlar=satirlar, genislik=[2.2, 2.2, 2.1, 3.1, 3.1], boyut=11.5,
             satir_h=1.0, kaynak="Sıklık verileri ilgili makalelerden aynen aktarılmıştır; tam künyeler bölümün çalışma slaytlarında ve kaynakçada.")
    s = tablo(prs, d, sayfa, bolum)
    # etiketi değiştir
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text == "TABLO":
            sh.text_frame.paragraphs[0].runs[0].text = "LİTERATÜR TARAMASI"; sh.fill.fore_color.rgb = ACCENT
    return s

def kaynakca(prs, pmcler, baslik, sayfa_bas):
    pmcler = list(dict.fromkeys(pmcler))
    n = 10; sayfa = sayfa_bas
    for i in range(0, len(pmcler), n):
        s = yeni_slayt(prs)
        title_bar(s, baslik + (f" ({i // n + 1})" if len(pmcler) > n else ""))
        items = []
        for p in pmcler[i:i + n]:
            e = K.makale(p)
            items.append(f"{e['authors']} {K.baslik_noktali(e)} ~{e['journal']}~ {e['year']}. doi:{e['doi']} – {p} – **{e['license_code']}**")
        bullets(s, 0.3, 1.0, 12.7, 5.95, items, size=9.5)
        footer(s, "Görseller PMC Open Access AWS deposundan (pmc-oa-opendata); lisans, md5 ve altyazılar sources/*.json içinde.", sayfa); fix_runs(s)
        sayfa += 1
    return sayfa
