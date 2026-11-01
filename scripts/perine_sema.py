# -*- coding: utf-8 -*-
"""Perine sunumu için yerel (düzenlenebilir) şemalar. Gerçek görüntü değildir; standart anatomi
bilgisinden çizilmiştir ve slaytta 'ŞEMA' olarak işaretlenir."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from deck_common import NAVY, ACCENT, GREY, LIGHT, WHITE, text, fix_runs, notes
import perine_slayt as P

KEMIK = RGBColor(0xD7, 0xC4, 0x9E); KAS = RGBColor(0xC6, 0x5B, 0x4F); FASYA = RGBColor(0x8E, 0xA9, 0xC8)
YAG = RGBColor(0xF6, 0xE3, 0x8C); BOS = RGBColor(0xEE, 0xF4, 0xFA); DERI = RGBColor(0xE8, 0xB9, 0x9A)
YESIL = RGBColor(0x2E, 0x7D, 0x32); MAVI = RGBColor(0x1F, 0x5F, 0x9E); SARI = RGBColor(0xE0, 0xA1, 0x00)
SEMA_NOTU = "Şema: bu projede standart anatomi bilgisinden çizilmiştir (gerçek görüntü değildir)."


def _cizgi(s, x1, y1, x2, y2, renk=NAVY, kalin=2.25, kesik=False):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = renk; ln.line.width = Pt(kalin)
    if kesik:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def _nokta(s, x, y, t, dx=0.12, dy=-0.18, w=2.6, hiza=PP_ALIGN.LEFT, boyut=14, renk=NAVY):
    r = 0.08
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = ACCENT; d.line.color.rgb = WHITE
    text(s, x + dx if hiza == PP_ALIGN.LEFT else x - dx - w, y + dy, w, 0.4, t, size=boyut, color=renk, bold=True, align=hiza)


def _kutu(s, x, y, w, h, t, renk, yazi=WHITE, boyut=14, sekil=MSO_SHAPE.ROUNDED_RECTANGLE, kalin=False):
    k = s.shapes.add_shape(sekil, Inches(x), Inches(y), Inches(w), Inches(h))
    k.fill.solid(); k.fill.fore_color.rgb = renk; k.line.color.rgb = WHITE
    tf = k.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = t if isinstance(t, list) else [t]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        import deck_common as D
        D._add_runs(p, line)
        for r in p.runs:
            r.font.size = Pt(boyut if i == 0 else boyut - 2); r.font.color.rgb = yazi; r.font.name = "Calibri"
            r.font.bold = (i == 0) or kalin
    return k


def _ok(s, x1, y1, x2, y2, renk=GREY, kalin=2.0):
    ln = _cizgi(s, x1, y1, x2, y2, renk, kalin)
    # ok başı: XML ile tailEnd
    from pptx.oxml.ns import qn
    lnx = ln.line._get_or_add_ln()
    te = lnx.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    lnx.append(te)
    return ln


def perine_eskenar(prs, d, sayfa):
    """Perinenin eşkenar dörtgen sınırları ve iki üçgen (litotomi pozisyonu, alttan bakış)."""
    s = P.yeni(prs)
    P.baslik(s, d["baslik"], d.get("alt"), "ŞEMA", d.get("bolum"))
    # köşeler (inç)
    S = (4.2, 1.75); TL = (2.1, 4.15); TR = (6.3, 4.15); C = (4.2, 6.45)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3), Inches(1.2), Inches(7.2), Inches(5.6))
    bg.fill.solid(); bg.fill.fore_color.rgb = BOS; bg.line.fill.background()
    # üçgen dolguları (freeform)
    ff = s.shapes.build_freeform(Inches(S[0]), Inches(S[1]))
    ff.add_line_segments([(Inches(TR[0]), Inches(TR[1])), (Inches(TL[0]), Inches(TL[1]))]); u = ff.convert_to_shape()
    u.fill.solid(); u.fill.fore_color.rgb = RGBColor(0xBB, 0xD6, 0xF0); u.line.fill.background()
    ff = s.shapes.build_freeform(Inches(TL[0]), Inches(TL[1]))
    ff.add_line_segments([(Inches(TR[0]), Inches(TR[1])), (Inches(C[0]), Inches(C[1]))]); a = ff.convert_to_shape()
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xF3, 0xD1, 0xC4); a.line.fill.background()
    for p1, p2, ad in ((S, TL, "R. ischiopubicus"), (S, TR, ""), (TL, C, "Lig. sacrotuberale"), (TR, C, "")):
        _cizgi(s, *p1, *p2, KEMIK if ad.startswith("R.") or (p1 == S) else RGBColor(0x8D, 0x6E, 0x63), 6)
    _cizgi(s, *TL, *TR, NAVY, 2.25, kesik=True)
    text(s, 0.9, 2.3, 2.3, 0.6, "Ramus ischiopubicus", size=13, color=RGBColor(0x79, 0x55, 0x48), bold=True, align=PP_ALIGN.RIGHT)
    text(s, 0.35, 5.2, 2.2, 0.6, "Lig. sacrotuberale", size=13, color=RGBColor(0x79, 0x55, 0x48), bold=True, align=PP_ALIGN.RIGHT)
    _nokta(s, *S, "Symphysis pubica (lig. pubicum inferius)", dx=0.15, dy=-0.45, w=3.4)
    _nokta(s, *TL, "Tuber ischiadicum", dx=0.12, dy=0.05, w=2.0, hiza=PP_ALIGN.RIGHT)
    _nokta(s, *TR, "Tuber ischiadicum", dx=0.12, dy=0.05, w=2.0)
    _nokta(s, *C, "Os coccygis (apex)", dx=0.15, dy=-0.2, w=2.4)
    text(s, 2.9, 3.25, 2.6, 0.7, ["Regio urogenitalis", "(trigonum urogenitale)"], size=14, color=MAVI, bold=True, align=PP_ALIGN.CENTER)
    text(s, 2.9, 4.55, 2.6, 0.7, ["Regio analis", "(trigonum anale)"], size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    text(s, 4.6, 3.88, 3.1, 0.3, "linea interischiadica", size=12, color=NAVY, italic=True)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.08), Inches(5.05), Inches(0.24), Inches(0.24)); o.fill.solid(); o.fill.fore_color.rgb = RGBColor(0x8D, 0x6E, 0x63)
    text(s, 4.35, 5.0, 1.2, 0.3, "anus", size=12, color=NAVY, italic=True)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.12), Inches(4.07), Inches(0.16), Inches(0.16)); o.fill.solid(); o.fill.fore_color.rgb = NAVY
    text(s, 3.5, 4.2, 1.6, 0.3, "corpus perineale", size=11, color=NAVY, italic=True, align=PP_ALIGN.CENTER)
    P.maddeler(s, 7.8, 1.3, 5.25, 5.6, d["metin"], size=d.get("boyut", 16))
    P.dipnot(s, SEMA_NOTU, sayfa); fix_runs(s)
    if d.get("not"): notes(s, d["not"])


def katman_yigini(prs, d, sayfa):
    """Regio urogenitalis'te yüzeyden derine katmanlar (koronal düzlemde basitleştirilmiş yığın)."""
    s = P.yeni(prs)
    P.baslik(s, d["baslik"], d.get("alt"), "ŞEMA", d.get("bolum"))
    K = [  # (ad, renk, yükseklik, tip)  yukarıdan aşağı: DERİN → YÜZEYEL (hasta litotomide; üst = kraniyal)
        ("Fascia superior diaphragmatis pelvis", FASYA, 0.22, "f"),
        ("M. levator ani (diaphragma pelvis)", KAS, 0.42, "k"),
        ("Fascia inferior diaphragmatis pelvis", FASYA, 0.22, "f"),
        ("Recessus anterior fossae ischioanalis (yağ)", YAG, 0.34, "y"),
        ("SPATIUM PERINEI PROFUNDUM – m. sphincter urethrae externus, m. compressor urethrae, m. sphincter urethrovaginalis, m. transversus perinei profundus, gl. bulbourethralis", BOS, 0.62, "b"),
        ("Membrana perinei", RGBColor(0x45, 0x5A, 0x64), 0.24, "m"),
        ("SPATIUM PERINEI SUPERFICIALE – crura & bulbus (penis/clitoridis, vestibuli), m. ischiocavernosus, m. bulbospongiosus, m. transversus perinei superficialis, gl. vestibularis major", BOS, 0.62, "b"),
        ("Fascia perinei (Gallaudet; kasları saran yatırım fasyası)", FASYA, 0.22, "f"),
        ("Stratum membranosum telae subcutaneae perinei (Colles fasyası)", RGBColor(0x5C, 0x7F, 0xA8), 0.26, "f"),
        ("Panniculus adiposus telae subcutaneae perinei", YAG, 0.34, "y"),
        ("Cutis", DERI, 0.24, "d"),
    ]
    y = 1.3; x = 0.4; w = 7.3
    text(s, x, y - 0.05, w, 0.3, "▲ kraniyal (derin) – cavitas pelvis", size=12, color=GREY, italic=True); y += 0.28
    for ad, renk, h, tip in K:
        k = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        k.fill.solid(); k.fill.fore_color.rgb = renk; k.line.color.rgb = WHITE
        tf = k.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.1); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = ad
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12.5 if tip == "b" else 13); r.font.name = "Calibri"; r.font.bold = tip in ("m", "k", "f")
                r.font.color.rgb = WHITE if tip in ("m", "k") or renk == RGBColor(0x5C, 0x7F, 0xA8) else NAVY
        y += h + 0.02
    text(s, x, y, w, 0.3, "▼ kaudal (yüzeyel) – deri", size=12, color=GREY, italic=True)
    P.maddeler(s, 8.0, 1.3, 5.05, 5.6, d["metin"], size=d.get("boyut", 15))
    P.dipnot(s, SEMA_NOTU + " Kalınlıklar ölçekli değildir.", sayfa); fix_runs(s)
    if d.get("not"): notes(s, d["not"])


def ekstravazasyon(prs, d, sayfa):
    """Fascia perinei superficialis (Colles) devamlılığı ve bulbar urethra rüptüründe idrar yayılımı."""
    s = P.yeni(prs)
    P.baslik(s, d["baslik"], d.get("alt"), "ŞEMA", d.get("bolum"))
    a = _kutu(s, 0.5, 1.4, 3.6, 0.95, ["Stratum membranosum (Scarpa)", "paries abdominis anterior"], MAVI, boyut=15)
    b = _kutu(s, 0.5, 3.0, 3.6, 0.95, ["Tunica dartos + fascia penis superficialis", "scrotum / penis"], MAVI, boyut=15)
    c = _kutu(s, 0.5, 4.6, 3.6, 0.95, ["Colles fasyası", "regio urogenitalis"], MAVI, boyut=15)
    for k1, k2 in ((a, b), (b, c)):
        _cizgi(s, 2.3, k1.top / 914400 + k1.height / 914400, 2.3, k2.top / 914400, MAVI, 4)
    t1 = _kutu(s, 4.6, 1.4, 3.3, 0.95, ["Fascia lata'ya tutunma", "lig. inguinale'nin ~1 parmak altı"], GREY, boyut=14)
    t2 = _kutu(s, 4.6, 4.6, 3.3, 0.95, ["Arka sınır: membrana perinei", "arka kenarına kaynaşma (m. transversus perinei superf.)"], GREY, boyut=13)
    t3 = _kutu(s, 4.6, 3.0, 3.3, 0.95, ["Lateral: rami ischiopubici", "(periost / crista)"], GREY, boyut=14)
    for k, kk in ((a, t1), (b, t3), (c, t2)):
        _cizgi(s, 4.1, kk.top / 914400 + 0.47, 4.6, kk.top / 914400 + 0.47, GREY, 2)
    r = _kutu(s, 8.4, 4.6, 4.5, 0.95, ["Ruptura urethrae (pars spongiosa / bulbaris)", "straddle travması – fascia penis profunda (Buck) yırtık"], ACCENT, boyut=14)
    _ok(s, 10.65, 4.6, 10.65, 3.95, ACCENT, 3)
    e = _kutu(s, 8.4, 2.95, 4.5, 1.0, ["İdrar spatium perinei superficiale'ye", "→ scrotum, penis, karın ön duvarı (Scarpa altı)"], RGBColor(0xC2, 0x6A, 0x00), boyut=14)
    _ok(s, 10.65, 2.95, 10.65, 2.35, ACCENT, 3)
    z = _kutu(s, 8.4, 1.4, 4.5, 0.95, ["Uyluğa ve anal üçgene GEÇMEZ", "(fascia lata ve membrana perinei arka kenarı engeller)"], NAVY, boyut=14)
    P.maddeler(s, 0.5, 5.75, 12.4, 1.2, d["metin"], size=d.get("boyut", 15), aralik=2)
    P.dipnot(s, SEMA_NOTU, sayfa); fix_runs(s)
    if d.get("not"): notes(s, d["not"])


def pudendal_agac(prs, d, sayfa):
    """N. pudendus ve a. pudenda interna: seyir ve dallar."""
    s = P.yeni(prs)
    P.baslik(s, d["baslik"], d.get("alt"), "ŞEMA", d.get("bolum"))
    SIN = SARI; ART = RGBColor(0xB7, 0x1C, 0x1C)
    # seyir şeridi
    adim = [("Plexus sacralis (S2–S4, ramus anterior)", 1.3), ("Foramen ischiadicum majus\n(infrapiriforme)", 2.2),
            ("Lig. sacrospinale / spina ischiadica\n(n. pudendus arterin medialinde)", 3.1),
            ("Foramen ischiadicum minus", 4.0), ("Canalis pudendalis (Alcock)\nfascia obturatoria duplikatürü", 4.9)]
    for t, y in adim:
        _kutu(s, 0.4, y, 4.2, 0.75, t.split("\n"), NAVY, boyut=13.5)
    for (_, y1), (_, y2) in zip(adim, adim[1:]):
        _ok(s, 2.5, y1 + 0.75, 2.5, y2, NAVY, 2.5)
    # dallar
    dal = [("N. rectalis inferior (nn. anales inf.)", "m. sphincter ani ext., perianal deri", 1.35),
           ("N. perinealis – r. superficialis", "nn. scrotales / labiales posteriores", 2.3),
           ("N. perinealis – r. profundus", "bulbospongiosus, ischiocavernosus, transversi, sphincter urethrae ext.", 3.25),
           ("N. dorsalis penis / clitoridis", "membrana perinei üstünden, lig. suspensorium altından → glans", 4.2)]
    for ad, hedef, y in dal:
        _kutu(s, 5.3, y, 3.6, 0.8, [ad, hedef], SIN, NAVY, boyut=13.5)
        _ok(s, 4.6, 5.27, 5.3, y + 0.4, SIN, 2)
    adal = [("A. rectalis inferior", 1.35), ("A. perinealis → aa. scrotales/labiales post.", 2.15),
            ("A. bulbi penis / vestibuli; a. urethralis", 2.95), ("A. profunda penis / clitoridis (crus)", 3.75),
            ("A. dorsalis penis / clitoridis", 4.55)]
    text(s, 9.3, 1.05, 3.8, 0.3, "A. pudenda interna (truncus anterior a. iliacae int.)", size=12, color=ART, bold=True)
    for ad, y in adal:
        _kutu(s, 9.3, y, 3.75, 0.66, ad, ART, boyut=13)
    P.maddeler(s, 0.4, 5.85, 12.6, 1.1, d["metin"], size=d.get("boyut", 14.5), aralik=2)
    P.dipnot(s, SEMA_NOTU, sayfa); fix_runs(s)
    if d.get("not"): notes(s, d["not"])
